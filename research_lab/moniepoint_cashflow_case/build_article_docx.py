from pathlib import Path
import re,json
from PIL import Image
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).resolve().parents[2];FOLDER=ROOT/'publications/linkedin/moniepoint_cashflow_2026-10-14';MD=FOLDER/'MONIEPOINT_CASHFLOW_UNDERWRITE_FOR_COLLECTION_LINKEDIN_ARTICLE.md'
text=MD.read_text(encoding='utf8')
# Replace tall full diagrams with the readable overview and direct full-panel links.
for i,name in enumerate(['MONIEPOINT_WORKFLOW_01_EVIDENCE_AND_CALCULATIONS','MONIEPOINT_WORKFLOW_02_APPROVAL_AND_EXECUTION','MONIEPOINT_WORKFLOW_03_MONITORING_AND_LEARNING'],7):
    text=re.sub(r'!\[[^\]]*\]\(charts/'+name+r'\.png\)\s*\*Figure '+str(i)+r'\.[^\n]*\*\s*','',text)
marker='[Complete governed workflow diagram](charts/MONIEPOINT_GOVERNED_AGENTIC_CREDIT_WORKFLOW.png).'
if '![Governed cash flow credit overview]' not in text:
    if marker not in text:
        boundary = '## Kenya’s DCP and data-protection setting'
        if text.count(boundary) != 1:
            raise ValueError('Expected one Kenya section as the workflow insertion boundary')
        text = text.replace(boundary, marker + '\n\n' + boundary, 1)
    text=text.replace(marker,'''![Governed cash flow credit overview](charts/MONIEPOINT_GOVERNED_WORKFLOW_OVERVIEW.png)

*Figure 7. The source, calculation, policy, approval, execution and monitoring stages form one governed workflow. The complete diagrams and readable Word supplement retain the detailed paths.*

### Read the evidence and calculation path

The first panel answers the underwriting question before it reaches a score: whose cash is this, which records describe the same movement, and what was actually knowable at review? The intake path binds the purpose and source permission to a case, account map and dated evidence snapshot. The case's bank and wallet split makes the reconciliation concrete: a missing wallet month reduces coverage, a duplicate shares an economic-event ID, and a later-known receipt classification remains outside verified operating cash at the earlier cutoff. The calculation branch then builds the cash bridge, descriptive checks, uncertainty and a challenger review. The current engine supplies the declared Normal cash posterior and dated calendar. The broader hierarchy and action-effect interfaces are developed in the five-article series and are labelled as research extensions in the diagram.

[Full evidence and calculations flowchart](charts/MONIEPOINT_WORKFLOW_01_EVIDENCE_AND_CALCULATIONS.png)

### Read the approval and execution path

The second panel shows how evidence becomes an authorised action. Reconciled information reaches versioned policy and eligibility gates, which determine the available terms and required review. The reviewer receives source coverage, the KES25,000 residual, the credible and predictive intervals, the dated cash minimum, the requested payment and the smaller proposal. The synthetic maker records a proposal; a different checker records approval before a payment adapter returns its execution receipt. A revised schedule follows the same separation. The diagram also shows expiring instructions and a current-conditions check for a production adapter: the evidence, permissions and applicable terms should still match the authorised instruction when execution occurs. These deployment controls extend the case's versioned approval and receipt records into the complete operating contract developed in Article 5.

[Full approval and execution flowchart](charts/MONIEPOINT_WORKFLOW_02_APPROVAL_AND_EXECUTION.png)

### Read the monitoring and learning path

The third panel begins with observed receipts and the approved payment calendar. A lower reported flow first branches into source quality, coverage and classification checks. In the feed-gap fixture, the missing channel routes to repair while the merchant's economic cash remains unchanged. In the delayed-settlement fixture, the evidence supports a customer review and a prospective schedule change. The next records distinguish the original due date, revised obligation, actual payment, borrower shortfall and lender funding cost. Mature outcomes can then support a registered evaluation or model update. Articles 1 and 2 develop experimental allocation and sequential decisions; Article 3 treats selective outcomes; Article 4 connects timing to funding; Article 5 preserves authority, change control and audit evidence. The complete flowchart therefore connects an implemented case with clearly identified interfaces for subsequent validated work.

[Full monitoring and learning flowchart](charts/MONIEPOINT_WORKFLOW_03_MONITORING_AND_LEARNING.png)

'''+marker)
MD.write_text(text,encoding='utf8')
manifest=json.loads((FOLDER/'charts/MONIEPOINT_WORKFLOW_WORD_MANIFEST.json').read_text())
doc=Document();sec=doc.sections[0]
def margins(sec,landscape=False):
    sec.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    sec.page_width=Inches(11.69 if landscape else 8.27);sec.page_height=Inches(8.27 if landscape else 11.69)
    sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.75)
margins(sec)
for name,size in [('Normal',11),('Title',22),('Heading 1',15),('Heading 2',12.5),('Caption',9)]:
    st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial');st.element.rPr.rFonts.set(qn('w:cs'),'Arial')
# Clear inherited theme fonts and template paragraph rules.
for st in doc.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    for fonts in st.element.iter(qn('w:rFonts')):
        for key in list(fonts.attrib):
            if key.endswith('Theme') or key.endswith('theme'):
                del fonts.attrib[key]
doc.styles['Normal'].paragraph_format.space_after=Pt(7);doc.styles['Normal'].paragraph_format.line_spacing=1.06
for name in ['Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.space_before=Pt(13);doc.styles[name].paragraph_format.space_after=Pt(6)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('Nevil Maloba | ');r.font.size=Pt(8)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)

def rich(p,s):
    for token in re.split(r'(\[[^\]]+\]\([^)]+\)|\*\*.*?\*\*|\*.*?\*)',s):
        if not token: continue
        link = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',token)
        if link:
            label,target=link.groups()
            if target.startswith('charts/'):
                target='https://github.com/TheAlchemistNerd/underwrite-for-collection/blob/main/publications/linkedin/moniepoint_cashflow_2026-10-14/'+target
            element=OxmlElement('w:hyperlink')
            element.set(qn('r:id'),p.part.relate_to(target,RT.HYPERLINK,is_external=True))
            run=OxmlElement('w:r');props=OxmlElement('w:rPr')
            color=OxmlElement('w:color');color.set(qn('w:val'),'0563C1');props.append(color)
            underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');props.append(underline)
            run.append(props);value=OxmlElement('w:t');value.text=label;run.append(value)
            element.append(run);p._p.append(element)
        else:
            r=p.add_run(token.strip('*'))
            if token.startswith('**'): r.bold=True
            elif token.startswith('*'): r.italic=True

def picture(path,max_width=6.55,max_height=8.5):
    width,height=Image.open(path).size;wi=min(max_width,max_height*width/height)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    r=p.add_run();shape=r.add_picture(str(path),width=Inches(wi));shape._inline.docPr.set('descr',path.stem.replace('_',' '))

for line in text.splitlines():
    line=line.strip()
    if not line or line.startswith('<!--'):continue
    if line.startswith('# '):rich(doc.add_paragraph(style='Title'),line[2:]);continue
    if line.startswith('## '):rich(doc.add_paragraph(style='Heading 1'),line[3:]);continue
    if line.startswith('### '):rich(doc.add_paragraph(style='Heading 2'),line[4:]);continue
    if line.startswith('!['):
        image_path=FOLDER/re.match(r'!\[[^\]]*\]\(([^)]+)\)',line).group(1)
        if not image_path.exists():raise FileNotFoundError(image_path)
        picture(image_path,max_height=7.8);continue
    if line.startswith('*Figure '):
        p=doc.add_paragraph(style='Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;rich(p,line.strip('*'));continue
    p=doc.add_paragraph();rich(p,line)
    if line.startswith('[') and re.match(r'\[\d+\]',line):
        p.paragraph_format.space_after=Pt(5)
        p.paragraph_format.line_spacing=1.0
        for r in p.runs: r.font.size=Pt(10)
# Detailed diagram supplement, one readable section per page.
sec=doc.add_section(WD_SECTION_START.NEW_PAGE);margins(sec)
doc.add_heading('Detailed governed workflow',level=1)
doc.add_paragraph('The following pages expand Figure 7 into readable continuation panels. Each panel retains whole nodes and overlapping connection bands. Full PNG, SVG and Mermaid files accompany the article for upload and reuse. The main article explains the business question, evidence, calculations, rules, human decisions and outcome records behind each path.')
doc.add_paragraph('The implemented synthetic case supplies source reconciliation, dated cash, a Normal cash posterior, decision gates, simulated maker–checker approvals, receipts and monitoring. The advanced experimental, causal, sequential-control and portfolio branches identify interfaces developed in the five-article series for subsequent validation and implementation.')
for index,item in enumerate(manifest['segments']):
    land=item['orientation']=='landscape'
    if index:
        sec=doc.add_section(WD_SECTION_START.NEW_PAGE)
    margins(sec,land)
    doc.add_heading(item['title'],level=2)
    picture(FOLDER/'charts'/item['image'],max_width=min(item['recommended_width_inches'],9.8 if land else 6.5),max_height=5.7 if land else (7.15 if index==0 else 8.8))
    p=doc.add_paragraph(item['caption'],style='Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
output=FOLDER/'MONIEPOINT_CASHFLOW_UNDERWRITE_FOR_COLLECTION_LINKEDIN_ARTICLE.docx';doc.save(output)
print(json.dumps({'docx':str(output),'images':len(doc.inline_shapes),'article_words':len(text.split('## References')[0].split())},indent=2))

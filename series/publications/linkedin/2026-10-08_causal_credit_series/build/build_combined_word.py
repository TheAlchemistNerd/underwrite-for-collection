"""Build the complete Word edition from the PDF pipeline's numbered math source."""
from pathlib import Path
import re,json,subprocess,hashlib
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from PIL import Image
base=Path(__file__).resolve().parent
root=base.parent
stem='UNDERWRITE_FOR_COLLECTION_FIVE_ARTICLE_RESEARCH_EDITION'
source=base/(stem+'_NUMBERED.md')
text=source.read_text(encoding='utf-8')
numbering=json.loads((base/'equation_numbering_report.json').read_text(encoding='utf-8'))
entries={e['label']:e for e in numbering['entries']}
aux=(base/(stem+'_2026-10-08.aux')).read_text(encoding='utf-8')
page_by_label={}
for line in aux.splitlines():
    m=re.match(r'\\newlabel\{(eq:[^}]+)\}\{\{(?:\{)?[^}]+\}(?:\})?\{([^}]+)\}',line)
    if m:page_by_label[m[1]]=m[2]
assert len(page_by_label)==len(entries),(len(page_by_label),len(entries))
anchor=lambda label:'word-'+label.replace(':','-')
seen=[]
def display(m):
    number,label,latex=m.groups(); seen.append(label)
    assert entries[label]['number']==number
    assert latex.strip()==entries[label]['latex'].strip()
    # Pandoc translates the same TeX mathematics into editable Word equations.
    return '$$\n'+latex.strip()+'\n$$\n\n[('+number+')]{#'+anchor(label)+'}\n'
text=re.sub(r'```\{=latex\}\s*\\begin\{equation\}\s*\\tag\{([^}]+)\}\\label\{([^}]+)\}\s*([\s\S]*?)\\end\{equation\}\s*```',display,text)
assert len(seen)==len(entries)
text=re.sub(r'\\eqref\{([^}]+)\}',lambda m:'[('+entries[m[1]]['number']+')](#'+anchor(m[1])+')',text)
text=re.sub(r'\\pageref\{([^}]+)\}',lambda m:page_by_label[m[1]],text)
names=['01_offer_experiment_and_governed_learning','02_temporal_graph_bandits_mcts_graphsvx','02_evidence_to_sequential_collection_control','03_reject_inference_evidence_path','04_treasury_data_to_alco_approval','04_treasury_constrained_control_loop','05_bank_controlled_agentic_ai_workflow']
captions=["Offer experiment and governed learning","Temporal graphs, bandits, MCTS and GraphSVX","Evidence-to-sequential-control pathway","Reject-inference evidence pathway","Treasury data to ALCO approval","Treasury constrained control loop","Bank-controlled agentic AI workflow"]
figure_numbers=["1.1","2.1","2.2","3.1","4.3","4.4","5.4"]
idx=iter(zip(names,figure_numbers,captions))
def diagram(m):
    name,number,caption=next(idx)
    pic=root/'linkedin_diagrams'/f'{name}.png'
    with Image.open(pic) as im: width=min(6.2,7.2*im.width/im.height)
    return '!['+caption+'](<'+pic.as_posix()+'>){width='+str(width)+'in}\n\n*Figure '+number+'. '+caption+'*'
text=re.sub(r'```mermaid\s*\n.*?```',diagram,text,flags=re.S)
text=text.replace('\\clearpage','').replace('\\newpage','')
# Equivalent TeX spelling supported by Pandoc's OMML converter.
text=re.sub(r'\{\\rm\s+([^{}]*)\}',lambda m:r'{\mathrm{'+m[1]+'}}',text)
work=base/'word';work.mkdir(exist_ok=True)
md=work/'COMBINED_NUMBERED_TEX_MATH.md';md.write_text(text,encoding='utf-8')
output=root/'publications'/(stem+'_2026-10-08.docx')
reference=work/'article_reference.docx'
pandoc=Path(r'C:\Users\Nevo\AppData\Local\Pandoc\pandoc.exe')
proc=subprocess.run([str(pandoc),str(md),'--from=markdown+autolink_bare_uris+tex_math_dollars-implicit_figures','--to=docx','--standalone','--toc','--toc-depth=2',f'--reference-doc={reference}','--output',str(output)],cwd=root,capture_output=True,text=True,encoding='utf-8')
(work/'combined_word_pandoc.log').write_text(proc.stderr,encoding='utf-8')
assert proc.returncode==0,proc.stderr
assert not re.search(r'Could not convert|Could not fetch',proc.stderr),proc.stderr
doc=Document(output)
doc.styles['Heading 1'].paragraph_format.page_break_before=True
for p in doc.paragraphs:
    if re.fullmatch(r'\(A?[1-5]\.\d+\)',p.text):
        p.alignment=2;p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.space_after=Pt(6)
    if p._p.xpath('.//m:oMathPara'):
        p.paragraph_format.keep_with_next=True
        p.paragraph_format.keep_together=True
# Keep each equation-register row together across Word page boundaries.
for table in doc.tables:
    for row in table.rows:
        props=row._tr.get_or_add_trPr()
        if not props.xpath('./w:cantSplit'):
            props.append(OxmlElement('w:cantSplit'))
# Core properties make the shared construction explicit.
doc.core_properties.author='Nevil Maloba'
doc.core_properties.comments='Complete companion to the independently compiled TeX PDF; same numbered mathematical source, articles, appendices and figures.'
doc.save(output)
with ZipFile(output) as z:
    xml=etree.fromstring(z.read('word/document.xml'))
    ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math','w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    display_count=len(xml.findall('.//m:oMathPara',ns))
    drawings=len(xml.findall('.//w:drawing',ns))
    assert display_count==len(entries),(display_count,len(entries))
    assert drawings==len(re.findall(r'!\[',text)),(drawings,len(re.findall(r'!\[',text)))
    alltext=' '.join(xml.xpath('//w:t/text()|//m:t/text()',namespaces=ns))
    for n in range(1,6):assert f'Appendix A{n}' in alltext
    assert '23343888' in alltext
report={'shared_numbered_source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'all_equation_math_preserved':True,'numbered_display_equations':display_count,'embedded_figures':drawings,'complete_articles_and_appendices':True,'pdf_page_numbers_reused_in_register':True,'output':str(output),'sha256':hashlib.sha256(output.read_bytes()).hexdigest()}
(work/'combined_word_checks.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
print(json.dumps(report,indent=2))

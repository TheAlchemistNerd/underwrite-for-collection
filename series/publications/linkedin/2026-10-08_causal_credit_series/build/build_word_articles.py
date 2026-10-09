"""Create LinkedIn Word articles with ordinary Unicode text and embedded charts."""
from pathlib import Path
import json
import re
import subprocess
import sys
from zipfile import ZipFile
from lxml import etree
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')
root=Path(__file__).resolve().parent.parent
out=root/'publications'/'word'
work=root/'build'/'word'
out.mkdir(parents=True,exist_ok=True)
work.mkdir(parents=True,exist_ok=True)
reference=work/'article_reference.docx'
ref=Document()
sec=ref.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.left_margin=sec.right_margin=Inches(.7)
sec.top_margin=sec.bottom_margin=Inches(.7)
sec.footer_distance=Inches(.3)
for name,size in [('Normal',11),('Title',22),('Subtitle',13),('Heading 1',16),('Heading 2',14),('Heading 3',12),('Caption',10)]:
    st=ref.styles[name]
    st.font.name='Calibri'; st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.line_spacing=1.1
    st.paragraph_format.widow_control=True
    if name.startswith('Heading'):
        st.font.bold=True
        st.paragraph_format.keep_with_next=True
        st.paragraph_format.space_before=Pt(12)
ref.styles['Title'].font.bold=True
for st in ref.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE')
footer._p.append(fld)
ref.save(reference)

names=[
 ('01_offer_experiment_and_governed_learning','Offer experiment and governed learning'),
 ('02_temporal_graph_bandits_mcts_graphsvx','Temporal graphs, bandits, MCTS and GraphSVX'),
 ('02_evidence_to_sequential_collection_control','Evidence-to-sequential-control pathway'),
 ('03_reject_inference_evidence_path','Reject-inference evidence pathway'),
 ('04_treasury_data_to_alco_approval','Treasury data to ALCO approval'),
 ('04_treasury_constrained_control_loop','Treasury constrained control loop'),
 ('05_bank_controlled_agentic_ai_workflow','Bank-controlled agentic AI workflow'),
]
sources=[next(p for p in root.glob(f'ARTICLE_{i}_*.md') if 'LINKEDIN_POST' not in p.name) for i in range(1,6)]
image_count=0
report=[]
for i,source in enumerate(sources,1):
    part=source.read_text(encoding='utf-8')
    title=part.strip().splitlines()[0].removeprefix('# ')
    body=part.strip().split('\n',1)[1]
    def diagram(match):
        global image_count
        name,caption=names[image_count]
        image_count+=1
        pic=root/'linkedin_diagrams'/f'{name}.png'
        with Image.open(pic) as im:
            width=min(6.4,7.2*im.width/im.height)
        return f'![Figure {image_count}. {caption}](<{pic.as_posix()}>){{width={width:.3f}in}}'
    body=re.sub(r'```mermaid\s*\n.*?```',diagram,body,flags=re.S)
    md=work/(source.stem+'_UNICODE.md')
    md.write_text('---\ntitle: '+json.dumps(title,ensure_ascii=False)+'\nauthor: Nevil Maloba\nlang: en\n---\n'+body,encoding='utf-8')
    docx=out/(source.stem+'.docx')
    args=['pandoc',str(md),'--from=markdown-tex_math_dollars-tex_math_single_backslash-tex_math_double_backslash+autolink_bare_uris','--to=docx','--standalone',f'--reference-doc={reference}','--output',str(docx)]
    result=subprocess.run(args,cwd=root,capture_output=True,text=True,encoding='utf-8')
    (work/f'article_{i}_pandoc.log').write_text(result.stderr,encoding='utf-8')
    if result.returncode: raise RuntimeError(result.stderr)
    if 'Could not convert' in result.stderr: raise RuntimeError(result.stderr)
    doc=Document(docx)
    for st in doc.styles:
        for border in list(st.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    for p in doc.paragraphs:
        for border in list(p._p.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
        # The characters stay literal Unicode when copied from Word to LinkedIn.
        if len(p.text)<230 and ('=' in p.text or '∼' in p.text) and not p.text.startswith('['):
            p.paragraph_format.space_before=Pt(7)
            p.paragraph_format.space_after=Pt(7)
            p.paragraph_format.keep_together=True
            for run in p.runs:
                run.font.name='Cambria Math'
        else:
            for run in p.runs:
                if any(c in run.text for c in '∼⫫∈⁎ₚ₀₁₂₃₄₅₆₇₈₉ᵢⱼₜ'):
                    run.font.name='Cambria Math'
        if p._p.findall('.//'+qn('w:drawing')):
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
    doc.core_properties.author='Nevil Maloba'
    doc.core_properties.title=title
    doc.save(docx)
    with ZipFile(docx) as z:
        xml=etree.fromstring(z.read('word/document.xml'))
        ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math','w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        item={'article':i,'file':str(docx),'editable_math':len(xml.findall('.//m:oMath',ns)),'display_equations':len(xml.findall('.//m:oMathPara',ns)),'diagrams':len(xml.findall('.//w:drawing',ns)),'bytes':docx.stat().st_size}
        assert item['editable_math']==0, 'LinkedIn Word export must use plain Unicode'
        assert item['diagrams']==[1,2,1,2,1][i-1]
        item['unicode_math']=True
        item['plain_text_characters']=sum(len(x) for x in xml.xpath('//w:t/text()',namespaces=ns))
        report.append(item)
    print(f'Article {i}: plain Unicode text; {item["diagrams"]} embedded diagrams')
(work/'word_export_checks.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
print('Created all five Unicode Word articles')

# One uploadable reading document retains the same ordinary Unicode text.
combined_parts=[]
for source in sources:
    prepared=(work/(source.stem+'_UNICODE.md')).read_text(encoding='utf-8')
    body=prepared.split('---\n',2)[2]
    title=source.read_text(encoding='utf-8').splitlines()[0]
    combined_parts.append(title+'\n'+body)
combined_md=work/'COMBINED_UNICODE.md'
combined_md.write_text('---\ntitle: "Underwrite for Collection Five Connected Research Articles"\nauthor: Nevil Maloba\nlang: en\n---\n\n'+'\n\n'.join(combined_parts),encoding='utf-8')
combined_docx=root/'publications/UNDERWRITE_FOR_COLLECTION_FIVE_ARTICLE_RESEARCH_EDITION_2026-10-08.docx'
subprocess.run(['pandoc',str(combined_md),'--from=markdown-tex_math_dollars-tex_math_single_backslash-tex_math_double_backslash+autolink_bare_uris','--to=docx','--standalone',f'--reference-doc={reference}','--output',str(combined_docx)],check=True,cwd=root)
doc=Document(combined_docx)
doc.styles['Heading 1'].paragraph_format.page_break_before=True
for st in doc.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for p in doc.paragraphs:
    for border in list(p._p.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    if len(p.text)<230 and ('=' in p.text or '∼' in p.text) and not p.text.startswith('['):
        p.paragraph_format.space_before=Pt(7); p.paragraph_format.space_after=Pt(7)
        p.paragraph_format.keep_together=True
        for run in p.runs: run.font.name='Cambria Math'
    else:
        for run in p.runs:
            if any(c in run.text for c in '∼⫫∈⁎ₚ₀₁₂₃₄₅₆₇₈₉ᵢⱼₜ'): run.font.name='Cambria Math'
    if p._p.findall('.//'+qn('w:drawing')):
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next=True
doc.core_properties.author='Nevil Maloba'
doc.save(combined_docx)
with ZipFile(combined_docx) as z:
    xml=etree.fromstring(z.read('word/document.xml'))
    assert len(xml.findall('.//m:oMath',ns))==0
    assert len(xml.findall('.//w:drawing',ns))==7
print('Created combined Unicode Word edition with all seven charts')
